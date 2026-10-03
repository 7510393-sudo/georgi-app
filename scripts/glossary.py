#!/usr/bin/env python3
"""Таблицы надписей en | ru | язык для автора: docs/translations/*.xlsx.

    python3 scripts/glossary.py
"""
import sys,re,os
sys.path.insert(0,'scripts')
import strings as S
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

names={"uk":"Українська","de":"Deutsch","fr":"Français","es":"Español","it":"Italiano","pt":"Português","nl":"Nederlands","ja":"日本語"}
want=S.keys()
order=sorted(want, key=lambda k:(k.strip(' \n·,.:“”«»…[]%@—').lower(), k))
yml=open('ios/project.yml',encoding='utf-8').read()
perm_en={}
for m in re.finditer(r'INFOPLIST_KEY_(NS\w+UsageDescription): "((?:[^"\\]|\\.)*)"', yml):
    perm_en.setdefault(m.group(1), m.group(2))
ru_perm=S.read_strings('ios/Chronotheca/ru.lproj/InfoPlist.strings')
perm_title={
 "NSMicrophoneUsageDescription":"Microphone / Микрофон","NSLocationWhenInUseUsageDescription":"Location / Местоположение",
 "NSCameraUsageDescription":"Camera / Камера","NSPhotoLibraryUsageDescription":"Photos / Фото",
 "NSCalendarsFullAccessUsageDescription":"Calendar / Календарь","NSCalendarsUsageDescription":"Calendar (older iOS) / Календарь (старые iOS)",
 "NSCalendarsWriteOnlyAccessUsageDescription":"Calendar, editing / Календарь, правка","NSSpeechRecognitionUsageDescription":"Speech recognition / Распознавание речи",
 "NSHealthShareUsageDescription":"Health, reading / Здоровье, чтение","NSHealthUpdateUsageDescription":"Health, writing / Здоровье, запись",
 "NSFaceIDUsageDescription":"Face ID",
}
head_font=Font(name='Arial',bold=True,color='FFFFFF',size=11)
fill=PatternFill('solid',start_color='2F4A6B')
body=Font(name='Arial',size=10.5)
wrap=Alignment(wrap_text=True,vertical='top')
thin=Border(bottom=Side(style='thin',color='DDDDDD'))
def show(s): return s.replace('\n','⏎\n') if s else s
def sheet(ws,title_cols,rows,widths):
    ws.append(title_cols)
    for c in ws[1]:
        c.font=head_font; c.fill=fill; c.alignment=Alignment(vertical='center',wrap_text=True)
    for r in rows:
        ws.append([show(x) for x in r])
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font=body; c.alignment=wrap; c.border=thin
    for i,w in enumerate(widths):
        ws.column_dimensions[chr(65+i)].width=w
    ws.freeze_panes='A2'
    ws.auto_filter.ref=ws.dimensions
for lang,lname in names.items():
    app=S.read_strings(f'ios/Chronotheca/{lang}.lproj/App.strings')
    perm=S.read_strings(f'ios/Chronotheca/{lang}.lproj/InfoPlist.strings')
    wb=Workbook()
    ws=wb.active; ws.title='Interface · Интерфейс'
    rows=[(k,want[k],app.get(k,'')) for k in order]
    sheet(ws,['English','Русский',lname],rows,[48,48,48])
    ws2=wb.create_sheet('Permissions · Разрешения')
    prow=[(perm_title[k],perm_en[k],ru_perm.get(k,''),perm.get(k,'')) for k in perm_title if k in perm_en]
    sheet(ws2,['Permission · Разрешение','English','Русский',lname],prow,[26,55,55,55])
    ws3=wb.create_sheet('About · О таблице')
    for line in ['Chronotheca — interface words and messages · надписи и сообщения приложения',
                 f'Languages · Языки: English, Русский, {lname}',
                 f'Interface lines · Строк интерфейса: {len(rows)}; permission questions · вопросов о разрешениях: {len(prow)}',
                 '⏎ — line break inside a message · перенос строки внутри сообщения',
                 '%@ — a value the app inserts (number, name, date) · вставка: число, название, дата',
                 'Some lines are pieces of one longer sentence; the app joins them · некоторые строки — куски одной фразы, приложение их склеивает',
                 'Source: ios/Chronotheca/*.lproj, checked 04.10.2026 · Источник: файлы переводов приложения, проверено 04.10.2026']:
        ws3.append([line])
    for r in ws3.iter_rows():
        for c in r: c.font=body
    ws3['A1'].font=Font(name='Arial',bold=True,size=12)
    ws3.column_dimensions['A'].width=110
    wb.save(f'docs/translations/Chronotheca-en-ru-{lang}.xlsx')
    print(lang,len(rows),len(prow))
